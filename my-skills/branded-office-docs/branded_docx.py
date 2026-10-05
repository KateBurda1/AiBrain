#!/usr/bin/env python3
"""Turn a markdown file into a KB&Co-branded Word doc.

Usage: python3 branded_docx.py [--letterhead] "<file.md>" [more.md ...]
Writes "<file>.docx" next to each .md, forces Franklin Gothic Book everywhere,
colors headings PMS 675 / PMS 214, and gives tables a PMS 675 header row with grey rules.

Two looks (Kate, 2026.09.28):
  --letterhead  Reports. Built on the real KB&Co Letterhead: KB logo at the top
                and the contact line at the bottom of page 1, content right under the logo.
  (default)     Plans and write-ups. Built on the KB&Co Word Template, with a
                cover page: the title and the details under it sit alone on
                page 1, and the content starts on page 2.
"""
import os
import subprocess
import sys

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt, RGBColor

BRAIN = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
TEMPLATE = os.path.join(BRAIN, "my-files (knowledge)", "Brand", "Kate & Co- Administrative",
                        "4. Word Template", "Word Template.docx")
LETTERHEAD = os.path.join(BRAIN, "my-files (knowledge)", "Brand", "Kate & Co- Administrative",
                          "9. Envelopes & Letterhead", "Letterhead.docx")
FONT = "Franklin Gothic Book"
PMS_675 = RGBColor(0xB5, 0x23, 0x72)
PMS_214 = RGBColor(0xD5, 0x10, 0x67)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_fonts(el):
    for rf in el.iter(qn("w:rFonts")):
        for a in list(rf.attrib):
            if a.endswith("Theme"):
                del rf.attrib[a]
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), FONT)


def cover_page(d):
    """Title and its details alone on page 1, content from page 2: a page
    break goes in front of the first section heading after the title."""
    paras = d.paragraphs
    if not paras:
        return
    paras[0].paragraph_format.space_before = Pt(180)
    for run in paras[0].runs:
        run.font.size = Pt(28)
    for p in paras[1:]:
        if p.style.name.startswith("Heading"):
            brk = p.insert_paragraph_before()
            brk.add_run().add_break(WD_BREAK.PAGE)
            # drop a horizontal rule left sitting alone at the foot of the cover
            prev = brk._p.getprevious()
            if prev is not None and prev.find(".//" + qn("w:pBdr")) is not None \
                    and not "".join(prev.itertext()).strip():
                prev.getparent().remove(prev)
            return


def brand(path, letterhead=False):
    d = docx.Document(path)
    if letterhead:
        # the letterhead's contact line sits inside the 1" bottom margin; clear it
        for sec in d.sections:
            sec.bottom_margin = Inches(1.6)
        h1 = d.styles["Heading 1"]
        h1.font.size = Pt(22)
        h1.font.all_caps = False
        # the letterhead's Heading 1 has a pink fill (swallows pink text) and a box
        for tag in ("w:shd", "w:pBdr"):
            for el in list(h1.element.iter(qn(tag))):
                el.getparent().remove(el)
    else:
        cover_page(d)
    for st in d.styles:
        if st.type in (1, 2):
            rpr = st.element.get_or_add_rPr()
            if rpr.find(qn("w:rFonts")) is None:
                rpr.append(OxmlElement("w:rFonts"))
    set_fonts(d.styles.element)
    for r in d.element.body.iter(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            r.insert(0, rpr)
        if rpr.find(qn("w:rFonts")) is None:
            rpr.insert(0, OxmlElement("w:rFonts"))
    set_fonts(d.element.body)
    for name, color in (("Title", PMS_675), ("Heading 1", PMS_675), ("Heading 2", PMS_214), ("Heading 3", PMS_214)):
        try:
            d.styles[name].font.color.rgb = color
        except KeyError:
            pass
    for t in d.tables:
        tblPr = t._tbl.tblPr
        for tag in ("w:tblBorders", "w:tblW"):
            for old in tblPr.findall(qn(tag)):
                tblPr.remove(old)
        borders = OxmlElement("w:tblBorders")
        for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{e}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "A0A1A2")
            borders.append(el)
        tblPr.append(borders)
        tw = OxmlElement("w:tblW")
        tw.set(qn("w:w"), "5000")
        tw.set(qn("w:type"), "pct")
        tblPr.append(tw)
        if not len(t.rows):
            continue
        for c in t.rows[0].cells:
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "B52372")
            c._tc.get_or_add_tcPr().append(shd)
            for p in c.paragraphs:
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = WHITE
    brand_bullets(d)
    d.save(path)


def brand_bullets(d):
    """Pandoc's bullets use a Symbol-font glyph that some viewers drop. Swap in
    a real bullet in Franklin Gothic, colored PMS 214, for every bullet level."""
    try:
        numbering = d.part.numbering_part.element
    except (KeyError, NotImplementedError):
        return
    glyphs = ["•", "–", "◦"]
    for lvl in numbering.iter(qn("w:lvl")):
        fmt = lvl.find(qn("w:numFmt"))
        if fmt is None or fmt.get(qn("w:val")) != "bullet":
            continue
        depth = int(lvl.get(qn("w:ilvl"), "0"))
        txt = lvl.find(qn("w:lvlText"))
        if txt is None:
            txt = OxmlElement("w:lvlText")
            lvl.append(txt)
        txt.set(qn("w:val"), glyphs[depth % len(glyphs)])
        rpr = lvl.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            lvl.append(rpr)
        for tag in ("w:rFonts", "w:color"):
            for old in rpr.findall(qn(tag)):
                rpr.remove(old)
        rf = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), FONT)
        rf.set(qn("w:hint"), "default")
        rpr.insert(0, rf)
        col = OxmlElement("w:color")
        col.set(qn("w:val"), "D51067")
        rpr.append(col)


args = sys.argv[1:]
use_letterhead = "--letterhead" in args
for md in [a for a in args if a != "--letterhead"]:
    out = os.path.splitext(md)[0] + ".docx"
    ref = LETTERHEAD if use_letterhead else TEMPLATE
    subprocess.run(["pandoc", md, f"--reference-doc={ref}", "-o", out], check=True)
    brand(out, letterhead=use_letterhead)
    print("branded:", out)
