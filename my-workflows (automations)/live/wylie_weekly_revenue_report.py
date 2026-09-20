#!/usr/bin/env python3
"""
Wylie Inn — Weekly Revenue Report builder.

Reads the source files Kate downloads every week and produces one branded,
formatted Word doc with the numbers already pulled out, deltas computed,
and outliers flagged — replacing the manual screenshot-and-paste step.

Usage:
    python3 wylie_weekly_revenue_report.py \\
      --str <STR .xlsx> --lighthouse <Lighthouse .xlsx> \\
      --booking-pdf <Booking.com print-to-PDF> \\
      --booking-promos <Booking.com ACTIVE promotions .xlsx> \\
      --synxis-csv <Synxis Inventory (Download) .csv> \\
      --expedia-promos-pdf <Expedia Promotions Report .pdf> \\
      --expedia-insights-png <Expedia Data & Insights screenshot> \\
      --expedia-travelads-png <Expedia TravelAds Performance screenshot> \\
      --expedia-chart-png <Expedia TravelAds performance chart screenshot> \\
      --out <output .docx>

Sources handled today:
  - CoStar/STR "Weekly STAR Report" (Glance + Segmentation Glance tabs)
  - Lighthouse rate export (Rates tab), joined with Synxis availability
  - Booking.com Sales statistics (print-to-PDF) + Promotions (.xlsx export)
  - Synxis Inventory (Download) CSV
  - Expedia Promotions Report (.pdf export) + Data Insights / TravelAds
    Performance (screenshots — no export exists for these two screens)

Not automatable:
  - Expedia Rate insights / Data & Insights hub / TravelAds Performance
    dashboard cards (confirmed no export — screenshots are the only option)
"""

import argparse
import csv
import datetime as dt
import re

import openpyxl
import pdfplumber
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DAY_NAMES = ["Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"]
EXCLUDED_MONTHS = {11}  # November left out per Kate — keeps the report to a tighter window

# Brand — Kate Burda & Company. Pulled from the real brand package:
# my-files (knowledge)/Brand/Kate & Co- Administrative/2. Style Guide & Fonts/
# and confirmed against the actual Word Template.docx in that same folder
# (Heading 1 = PMS 675, Subtitle/accent = PMS 214, Heading 2 = the dark gray,
# body font Franklin Gothic Book throughout). Green/red are a standard
# data-viz convention for above/below comp, not brand colors.
BRAND_PRIMARY_PINK = "D51067"    # PMS 214
BRAND_SECONDARY_PINK = "B52372"  # PMS 675
BRAND_DARK_GRAY = "404040"
BRAND_FONT = "Franklin Gothic Book"
GOOD_GREEN = "2E7D32"
BAD_RED = "C62828"
HEADER_TEXT = "FFFFFF"

BRAND_LOGO_PATH = (
    "/Users/kate/Library/CloudStorage/OneDrive-turningpointglobal.net/My-AI-Brain/"
    "My AI Brain/my-workflows (automations)/live/assets/kb_co_logo.png"
)


def pct(x):
    if x is None:
        return "--"
    return f"{x * 100:.1f}%"


def money(x):
    if x is None:
        return "--"
    return f"${x:,.0f}"


def idx(x):
    if x is None:
        return "--"
    return f"{x:.1f}"


def parse_str_glance(path):
    """Pull the 'This week' block out of the STR Glance tab: Occ/ADR/RevPAR
    and Index (MPI/ARI/RGI) for My Property vs Comp Set, by day."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Glance"]

    header_row = None
    for r in range(1, 15):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, 25)]
        if "Sunday" in row_vals:
            header_row = r
            break
    if header_row is None:
        raise ValueError("Could not find day header row in Glance tab")

    day_cols = {}
    for c in range(1, 25):
        v = ws.cell(row=header_row, column=c).value
        if v in DAY_NAMES:
            day_cols[v] = c

    metrics = {}
    r = header_row
    current_metric = None
    while r < header_row + 15:
        r += 1
        label = ws.cell(row=r, column=2).value
        sublabel = ws.cell(row=r, column=4).value
        if label in ("Occupancy", "ADR", "RevPAR"):
            current_metric = label
            metrics.setdefault(current_metric, {})
        if current_metric and sublabel in ("My Property", "Comp Set", "Index (MPI)", "Index (ARI)", "Index (RGI)"):
            row_data = {}
            for day, col in day_cols.items():
                row_data[day] = ws.cell(row=r, column=col).value
            metrics[current_metric][sublabel] = row_data
        if label == "Running 28 Day":
            break

    return metrics


def parse_str_segmentation(path):
    """Pull the 'This week' segment block: Transient/Group/Contract x
    Occ/ADR/RevPAR, My Property vs Comp set vs Index."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Segmentation Glance"]

    header_row = None
    for r in range(1, 12):
        vals = [ws.cell(row=r, column=c).value for c in range(1, 15)]
        if "Transient" in vals:
            header_row = r
            break
    if header_row is None:
        raise ValueError("Could not find segment header row")

    seg_cols = {}
    for c in range(1, 15):
        v = ws.cell(row=header_row, column=c).value
        if v in ("Transient", "Group", "Contract"):
            seg_cols[v] = c

    segments = {}
    current_metric = None
    r = header_row + 1
    while r < header_row + 12:
        r += 1
        label = ws.cell(row=r, column=2).value
        if label in ("Occupancy", "ADR", "RevPAR"):
            current_metric = label
            segments.setdefault(current_metric, {})
        if current_metric:
            row_label = ws.cell(row=r, column=list(seg_cols.values())[0]).value
            if row_label in ("My Property", "Comp set", "Index (MPI)", "Index (ARI)", "Index (RGI)"):
                for seg, col in seg_cols.items():
                    val = ws.cell(row=r, column=col + 1).value
                    segments[current_metric].setdefault(row_label, {})[seg] = val
        if label == "Running 28 Days":
            break

    return segments


def parse_lighthouse_rates(path):
    """Pull the Rates tab: date, day, my rate, each comp hotel's rate."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Rates"]

    header_row = None
    day_col = None
    for r in range(1, 10):
        for c in range(1, 6):
            if ws.cell(row=r, column=c).value == "Day" and ws.cell(row=r, column=c + 1).value == "Date":
                header_row, day_col = r, c
                break
        if header_row:
            break
    if header_row is None:
        raise ValueError("Could not find 'Day'/'Date' header row in Rates tab")

    headers = []
    c = day_col
    while True:
        v = ws.cell(row=header_row, column=c).value
        if v is None and c > day_col + 1:
            break
        headers.append(v)
        c += 1

    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        date_val = ws.cell(row=r, column=day_col + 1).value
        if date_val is None:
            continue
        if isinstance(date_val, dt.datetime) and date_val.month in EXCLUDED_MONTHS:
            continue
        row = {}
        for offset, h in enumerate(headers):
            if h:
                row[h] = ws.cell(row=r, column=day_col + offset).value
        rows.append(row)

    return headers, rows


def _dedupe_chars(chars):
    """Booking.com's PDF export renders bold text as two near-identical
    overlapping glyphs per character (a faux-bold artifact). Plain
    extract_text() interleaves them into 'RRoooomm'-style garbage, so pull
    from page.chars directly and drop the near-duplicate at (almost) the
    same position."""
    seen = set()
    out = []
    for ch in chars:
        key = (ch["text"], round(ch["x0"], 1), round(ch["top"], 1))
        if key in seen:
            continue
        seen.add(key)
        out.append(ch)
    return out


def _chars_to_rows(chars, gap=2):
    from collections import defaultdict
    rows = defaultdict(list)
    for ch in chars:
        rows[round(ch["top"], 1)].append(ch)
    lines = []
    for top in sorted(rows.keys()):
        row_chars = sorted(rows[top], key=lambda c: c["x0"])
        words, cur, last_x1 = [], "", None
        for c in row_chars:
            if last_x1 is not None and c["x0"] - last_x1 > gap:
                words.append(cur)
                cur = ""
            cur += c["text"]
            last_x1 = c["x1"]
        if cur:
            words.append(cur)
        lines.append((top, words))
    return lines


def parse_booking_pdf(path):
    """Parse the Booking.com 'Sales statistics' print-to-PDF: monthly room
    nights / revenue / ADR vs. last year (page 1), and the room/rate-plan
    breakdown (page 2)."""
    with pdfplumber.open(path) as pdf:
        page0_text = pdf.pages[0].extract_text()
        page1_chars = _dedupe_chars(pdf.pages[1].chars)

    lines = [l for l in page0_text.split("\n") if "Claude is active" not in l]

    months_row_idx = next(
        i for i, l in enumerate(lines) if re.match(r"^\d{4}(\s+\d{4}){3,8}\s+Total$", l)
    )
    months = [m for m in lines[months_row_idx - 1].split()]

    row_keys = [
        "room_nights", "room_nights_ly", "room_nights_pct_diff",
        "revenue", "revenue_ly", "revenue_pct_diff",
        "adr", "adr_ly", "adr_pct_diff",
    ]
    monthly = {}
    for offset, key in enumerate(row_keys, start=1):
        line = lines[months_row_idx + offset]
        m = re.match(r"^(.*?)\s+([+\-\d].*)$", line)
        if not m:
            continue
        monthly[key] = m.group(2).split()

    row_lines = _chars_to_rows(page1_chars)
    breakdown = []
    for top, words in row_lines:
        if len(words) == 4 and words[0] not in ("Room/Rate plan",):
            label, room_nights, revenue, adr = words
            breakdown.append({
                "label": label,
                "room_nights": room_nights,
                "revenue": revenue,
                "adr": adr,
                "is_total": label == "Total",
            })

    return {"months": months, "monthly": monthly, "breakdown": breakdown}


def parse_booking_promotions_xlsx(path):
    """Parse Booking.com's ACTIVE promotions .xlsx export."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Promotions"]
    rows = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=1).value
        bookings = ws.cell(row=r, column=5).value
        if not name or bookings in (None, ""):
            continue
        rows.append({
            "name": name,
            "discount": ws.cell(row=r, column=2).value,
            "bookings": bookings,
            "room_nights": ws.cell(row=r, column=6).value,
            "adr": ws.cell(row=r, column=7).value,
            "revenue": ws.cell(row=r, column=8).value,
        })
    return rows


def parse_expedia_promotions_pdf(path):
    """Parse Expedia's Promotions Report .pdf export — clean text, no
    faux-bold doubling issue like the Booking.com export has."""
    with pdfplumber.open(path) as pdf:
        text = "\n".join(p.extract_text() for p in pdf.pages)

    lines = text.split("\n")

    rows = []
    matched_lines = set()
    # Top-level promo rows start with a name and end in a $ gross value, a
    # $ net value, two integers (bookings, room nights), then dates/status.
    pattern = re.compile(
        r"^(?P<name>.+?)\s+(?P<offer>[\d/]+%)\s+(?:Standard|Campaign)\s+\S.*?"
        r"\$(?P<gross>[\d,\.]+)\s+\$(?P<net>[\d,\.]+)\s+(?P<bookings>\d+)\s+(?P<room_nights>\d+)\s+"
        r"(?P<stay_start>\S+)\s+(?P<stay_end>\S.*?)\s+(?P<book_start>\S+)\s+(?P<book_end>\S.*?)\s+"
        r"(?:\d+\s+)?(?P<status>Active|Inactive|Expired)$"
    )
    for i, line in enumerate(lines):
        m = pattern.match(line.strip())
        if m:
            rows.append({
                "name": m.group("name"),
                "offer": m.group("offer"),
                "gross": m.group("gross"),
                "net": m.group("net"),
                "bookings": m.group("bookings"),
                "room_nights": m.group("room_nights"),
                "status": m.group("status"),
            })
            matched_lines.add(i)

    # Fallback for rows whose "Offer" column wrapped onto the next line
    # (breaks the strict pattern above, e.g. "10 to" / "20% only" split
    # across two lines) — loosen the offer requirement and pull whatever
    # revenue/booking figures are present.
    fallback = re.compile(
        r"^(?P<name>.+?)\s+(?:Standard|Campaign)\s+\S.*?"
        r"\$(?P<gross>[\d,\.]+)\s+\$(?P<net>[\d,\.]+)\s+(?P<bookings>\d+)\s+(?P<room_nights>\d+)\s+"
        r"(?P<stay_start>\S+)\s+(?P<stay_end>\S.*?)\s+(?P<book_start>\S+)\s+(?P<book_end>\S.*?)\s+"
        r"(?:\d+\s+)?(?P<status>Active|Inactive|Expired)$"
    )
    for i, line in enumerate(lines):
        if i in matched_lines:
            continue
        m = fallback.match(line.strip())
        if m:
            name = re.sub(r"\s+\d+\s+to$", "", m.group("name"))  # strip a wrapped "N to" fragment
            rows.append({
                "name": name,
                "offer": "--",
                "gross": m.group("gross"),
                "net": m.group("net"),
                "bookings": m.group("bookings"),
                "room_nights": m.group("room_nights"),
                "status": m.group("status"),
            })

    return rows


def parse_synxis_inventory(path):
    """Parse the Synxis 'Inventory (Download)' CSV: total available/sold
    rooms per date, summed across room types."""
    from collections import defaultdict
    totals = defaultdict(lambda: {"avail": 0, "sold": 0})
    with open(path, encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            if not row.get("Hotel_Id") or not row.get("Cal_Dt"):
                continue
            date = dt.datetime.strptime(row["Cal_Dt"], "%m/%d/%Y %I:%M:%S %p").date()
            if date.month in EXCLUDED_MONTHS:
                continue
            totals[date]["avail"] += int(row["Avail_Qty"])
            totals[date]["sold"] += int(row["Sell_Qty"])
    return totals


# --- docx helpers -----------------------------------------------------

def _set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _set_run_font(run, size=None):
    run.font.name = BRAND_FONT
    # Word keys the East Asian/complex-script font slots separately — set
    # them too or some builds silently fall back to Calibri for parts of it.
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), BRAND_FONT)
    if size:
        run.font.size = Pt(size)


def _set_cell_text(cell, text, bold=False, color=None, size=None):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run("" if text is None else str(text))
    run.bold = bold
    _set_run_font(run, size=size or 10)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    color = {1: BRAND_SECONDARY_PINK}.get(level, BRAND_DARK_GRAY)
    for run in h.runs:
        run.font.color.rgb = RGBColor.from_string(color)
        _set_run_font(run, size={1: 16, 2: 13}.get(level, 12))
        if level == 2:
            run.bold = True
    return h


def add_table(doc, header, rows, shade_col=None, shade_fn=None):
    """shade_col: index of a column whose cells get colored by shade_fn(value).
    shade_fn takes the raw cell string and returns a hex color or None."""
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(header):
        _set_cell_text(hdr_cells[i], h, bold=True, color=HEADER_TEXT)
        _set_cell_shading(hdr_cells[i], BRAND_PRIMARY_PINK)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            _set_cell_text(cells[i], val)
            if shade_col is not None and i == shade_col and shade_fn:
                color = shade_fn(val)
                if color:
                    _set_cell_shading(cells[i], color)
    return table


def add_image(doc, image_path, width_inches=6.3):
    doc.add_picture(image_path, width=Inches(width_inches))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def set_document_font(doc):
    """Franklin Gothic Book as the base body font, matching the real KB Co
    Word Template (my-files (knowledge)/Brand/.../4. Word Template/)."""
    style = doc.styles["Normal"]
    style.font.name = BRAND_FONT
    style.font.size = Pt(10)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), BRAND_FONT)


def add_letterhead(doc):
    """KB Co logo in the header, right-aligned, matching the real template."""
    import os
    if not os.path.exists(BRAND_LOGO_PATH):
        return
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    run.add_picture(BRAND_LOGO_PATH, width=Inches(1.1))


def index_shade(val):
    try:
        v = float(val)
    except (TypeError, ValueError):
        return None
    return GOOD_GREEN if v >= 100 else BAD_RED


def build_report(lighthouse_path, out_path, str_path=None, booking_pdf_path=None,
                  synxis_csv_path=None, booking_promos_path=None,
                  expedia_promos_pdf_path=None, expedia_insights_png=None,
                  expedia_travelads_png=None, expedia_chart_png=None):
    glance = parse_str_glance(str_path) if str_path else None
    segmentation = parse_str_segmentation(str_path) if str_path else None
    lh_headers, lh_rows = parse_lighthouse_rates(lighthouse_path)
    booking = parse_booking_pdf(booking_pdf_path) if booking_pdf_path else None
    synxis = parse_synxis_inventory(synxis_csv_path) if synxis_csv_path else None
    booking_promos = parse_booking_promotions_xlsx(booking_promos_path) if booking_promos_path else None
    expedia_promos = parse_expedia_promotions_pdf(expedia_promos_pdf_path) if expedia_promos_pdf_path else None

    doc = Document()
    set_document_font(doc)
    add_letterhead(doc)

    title = doc.add_heading("The Wylie Inn — Weekly Revenue Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = RGBColor.from_string(BRAND_SECONDARY_PINK)
        _set_run_font(run, size=22)
    sub = doc.add_paragraph(f"Generated {dt.date.today().strftime('%B %d, %Y')}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in sub.runs:
        run.font.color.rgb = RGBColor.from_string(BRAND_PRIMARY_PINK)
        _set_run_font(run, size=11)
        run.italic = True

    if str_path:
        # --- STR / CoStar section ---
        add_heading(doc, "Market Performance vs. Comp Set (CoStar / STR)", level=1)
        days = [d for d in DAY_NAMES if d in glance.get("Occupancy", {}).get("My Property", {})]

        for metric, index_label in [("Occupancy", "Index (MPI)"), ("ADR", "Index (ARI)"), ("RevPAR", "Index (RGI)")]:
            add_heading(doc, metric, level=2)
            header = ["Day", "My Property", "Comp Set", index_label]
            rows = []
            for d in days:
                my_val = glance[metric]["My Property"].get(d)
                comp_val = glance[metric]["Comp Set"].get(d)
                index_val = glance[metric][index_label].get(d)
                fmt = pct if metric == "Occupancy" else money
                rows.append([d, fmt(my_val), fmt(comp_val), idx(index_val)])
            add_table(doc, header, rows, shade_col=3, shade_fn=index_shade)

            losing = [d for d in days if (glance[metric][index_label].get(d) or 100) < 100]
            if losing:
                p = doc.add_paragraph()
                p.add_run(f"Below comp set on {metric}: ").bold = True
                p.add_run(", ".join(losing))

        add_heading(doc, "Segmentation (This Week)", level=2)
        for metric in ["Occupancy", "ADR", "RevPAR"]:
            seg_data = segmentation.get(metric, {})
            header = ["Segment", "My Property", "Comp Set", "Index"]
            rows = []
            for seg in ["Transient", "Group", "Contract"]:
                my_val = seg_data.get("My Property", {}).get(seg)
                comp_val = seg_data.get("Comp set", {}).get(seg)
                index_key = [k for k in seg_data if k.startswith("Index")]
                index_val = seg_data.get(index_key[0], {}).get(seg) if index_key else None
                fmt = pct if metric == "Occupancy" else money
                rows.append([seg, fmt(my_val), fmt(comp_val), idx(index_val)])
            p = doc.add_paragraph()
            p.add_run(metric).bold = True
            add_table(doc, header, rows, shade_col=3, shade_fn=index_shade)

    # --- Lighthouse section (joined with Synxis availability, if provided) ---
    add_heading(doc, "Rate Shop vs. Comp Set (Lighthouse)", level=1)
    comp_names = [h for h in lh_headers if h and h not in ("Day", "Date")][1:]
    my_hotel_col = lh_headers[2] if len(lh_headers) > 2 else "My Hotel"

    header = ["Day", "Date"] + (["Wylie Avail/Sold"] if synxis else []) + [my_hotel_col] + comp_names
    rows = []
    for r in lh_rows:
        date_val = r.get("Date")
        date_str = date_val.strftime("%m/%d") if isinstance(date_val, dt.datetime) else str(date_val)
        row = [r.get("Day"), date_str]
        if synxis:
            key = date_val.date() if isinstance(date_val, dt.datetime) else None
            inv = synxis.get(key) if key else None
            row.append(f"{inv['avail']}/{inv['sold']}" if inv else "--")
        row.append(r.get(my_hotel_col))
        for c in comp_names:
            row.append(r.get(c))
        rows.append(row)
    add_table(doc, header, rows)

    if synxis:
        underpriced = []
        for r in lh_rows:
            date_val = r.get("Date")
            key = date_val.date() if isinstance(date_val, dt.datetime) else None
            inv = synxis.get(key) if key else None
            if not inv or inv["avail"] > 0:
                continue
            comp_open = any(isinstance(r.get(c), (int, float)) for c in comp_names)
            if comp_open:
                date_str = date_val.strftime("%m/%d") if isinstance(date_val, dt.datetime) else str(date_val)
                underpriced.append(date_str)
        if underpriced:
            p = doc.add_paragraph()
            p.add_run("Wylie sold out while comp set still has availability: ").bold = True
            p.add_run(", ".join(underpriced))

    outliers = []
    for r in lh_rows:
        my_rate = r.get(my_hotel_col)
        if not isinstance(my_rate, (int, float)):
            continue
        comp_rates = [r.get(c) for c in comp_names if isinstance(r.get(c), (int, float))]
        if comp_rates and my_rate < min(comp_rates):
            date_val = r.get("Date")
            date_str = date_val.strftime("%m/%d") if isinstance(date_val, dt.datetime) else str(date_val)
            outliers.append(date_str)
    if outliers:
        p = doc.add_paragraph()
        p.add_run("Priced below entire comp set: ").bold = True
        p.add_run(", ".join(outliers))

    # --- Booking.com section ---
    if booking:
        add_heading(doc, "Booking.com — Your Production", level=1)
        months = booking["months"] + ["Total"]

        header = ["Metric"] + months
        rows = [
            ["Room nights"] + booking["monthly"].get("room_nights", []),
            ["Room nights (last year)"] + booking["monthly"].get("room_nights_ly", []),
            ["% difference"] + booking["monthly"].get("room_nights_pct_diff", []),
            ["Revenue (USD)"] + booking["monthly"].get("revenue", []),
            ["Revenue (last year)"] + booking["monthly"].get("revenue_ly", []),
            ["% difference"] + booking["monthly"].get("revenue_pct_diff", []),
            ["ADR (USD)"] + booking["monthly"].get("adr", []),
            ["ADR (last year)"] + booking["monthly"].get("adr_ly", []),
            ["% difference"] + booking["monthly"].get("adr_pct_diff", []),
        ]
        add_table(doc, header, rows)

        behind = []
        for month, diff in zip(months, booking["monthly"].get("room_nights_pct_diff", [])):
            if diff.startswith("-"):
                behind.append(f"{month} ({diff})")
        if behind:
            p = doc.add_paragraph()
            p.add_run("Room nights behind last year: ").bold = True
            p.add_run(", ".join(behind))

        add_heading(doc, "Breakdown by Room and Rate Plan", level=2)
        header = ["Room / Rate Plan", "Room Nights", "Revenue (USD)", "ADR (USD)"]
        rows = [[r["label"], r["room_nights"], r["revenue"], r["adr"]] for r in booking["breakdown"]]
        add_table(doc, header, rows)

    if booking_promos:
        add_heading(doc, "Booking.com — Active Promotions", level=2)
        header = ["Promotion", "Discount", "Bookings", "Room Nights", "ADR", "Revenue"]
        rows = [[p["name"], p["discount"], p["bookings"], p["room_nights"], p["adr"], p["revenue"]]
                for p in booking_promos]
        add_table(doc, header, rows)

    # --- Expedia section ---
    add_heading(doc, "Expedia", level=1)

    if expedia_promos:
        add_heading(doc, "Active Promotions", level=2)
        header = ["Promotion", "Offer", "Gross Booking Value", "Net Revenue", "Bookings", "Room Nights", "Status"]
        rows = [[p["name"], p["offer"], f"${p['gross']}", f"${p['net']}", p["bookings"], p["room_nights"], p["status"]]
                for p in expedia_promos]
        add_table(doc, header, rows)

    if expedia_insights_png:
        add_heading(doc, "Data & Insights", level=2)
        doc.add_paragraph("No file export exists for this screen — captured as a screenshot.")
        add_image(doc, expedia_insights_png)

    if expedia_travelads_png:
        add_heading(doc, "TravelAds Performance", level=2)
        doc.add_paragraph("No file export exists for this screen — captured as a screenshot.")
        add_image(doc, expedia_travelads_png)

    if expedia_chart_png:
        add_image(doc, expedia_chart_png)

    if not any([expedia_promos, expedia_insights_png, expedia_travelads_png]):
        doc.add_paragraph(
            "No Expedia files or screenshots provided this run — pull the "
            "Promotions Report (.pdf, has a Download button) and screenshot "
            "the Data & Insights hub and TravelAds Performance dashboard by "
            "hand for this week's call."
        )

    add_heading(doc, "For the call", level=1)
    if synxis:
        doc.add_paragraph(
            "Wylie Avail/Sold is joined into the rate table above. Days "
            "flagged \"sold out while comp set still has availability\" are "
            "candidates to hold or raise rate; days where the comp set is "
            "selling out and Wylie isn't may be priced too high."
        )
    else:
        doc.add_paragraph(
            "No Synxis availability file provided this run — cross-check the "
            "rate table above against Wylie's Synxis availability (Reports > "
            "Inventory (Download)) by hand for this week's call."
        )

    doc.save(out_path)
    print(f"Saved: {out_path}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--str", dest="str_path", default=None,
                         help="Path to the STR Weekly STAR Report .xlsx (optional — skip for the default Lighthouse + Synxis-only run)")
    parser.add_argument("--lighthouse", dest="lighthouse_path", required=True, help="Path to the Lighthouse rate export .xlsx")
    parser.add_argument("--out", dest="out_path", required=True, help="Path to save the generated .docx")
    parser.add_argument("--booking-pdf", dest="booking_pdf_path", default=None,
                         help="Path to the Booking.com Sales statistics print-to-PDF (optional)")
    parser.add_argument("--booking-promos", dest="booking_promos_path", default=None,
                         help="Path to Booking.com's ACTIVE promotions .xlsx export (optional)")
    parser.add_argument("--synxis-csv", dest="synxis_csv_path", default=None,
                         help="Path to the Synxis Inventory (Download) report CSV (optional)")
    parser.add_argument("--expedia-promos-pdf", dest="expedia_promos_pdf_path", default=None,
                         help="Path to Expedia's Promotions Report .pdf export (optional)")
    parser.add_argument("--expedia-insights-png", dest="expedia_insights_png", default=None,
                         help="Screenshot of Expedia's Data & Insights hub (optional)")
    parser.add_argument("--expedia-travelads-png", dest="expedia_travelads_png", default=None,
                         help="Screenshot of Expedia's TravelAds Performance dashboard (optional)")
    parser.add_argument("--expedia-chart-png", dest="expedia_chart_png", default=None,
                         help="Screenshot of the TravelAds performance chart (optional)")
    args = parser.parse_args()
    build_report(
        args.lighthouse_path, args.out_path,
        str_path=args.str_path,
        booking_pdf_path=args.booking_pdf_path,
        synxis_csv_path=args.synxis_csv_path,
        booking_promos_path=args.booking_promos_path,
        expedia_promos_pdf_path=args.expedia_promos_pdf_path,
        expedia_insights_png=args.expedia_insights_png,
        expedia_travelads_png=args.expedia_travelads_png,
        expedia_chart_png=args.expedia_chart_png,
    )
